#!/usr/bin/env python3
# 從「Video 營收大表」重建 revenue.json。
# 用法：python3 revenue_import.py <輸入檔>
#   輸入檔可以是：
#   (a) Google Drive download_file_content 的 JSON 結果檔（content 欄位為 base64 xlsx）
#   (b) 直接的 .xlsx 檔
# 解析「營收大表(Video)-2026」分頁：欄位以第一列標題文字定位（來源表曾插入新欄導致欄位位移），
# 需要成員、合作名稱、產品項目、走期開始、成員收入(<80萬 5/5分)、成員收入(>80萬 6/4分)、成交價(未稅)。
import base64, json, os, sys, tempfile
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
SHEET = "營收大表(Video)-2026"
MAP = {"Rio": "rio", "Andrew": "andrew", "Ellie": "ellie", "YY": "yy", "梅伯": "meber",
       "Leo": "leo", "Nakaw": "nakaw", "Roy": "roy", "Royce": "royce",
       "主委": "roy", "吳至晟": "roy", "Lauren": "lauren", "阿根": "agen", "Anthony": "agen"}

src = sys.argv[1]
if src.endswith(".xlsx"):
    xlsx_path = src
else:
    data = json.load(open(src))
    xlsx_path = os.path.join(tempfile.gettempdir(), "revenue_export.xlsx")
    with open(xlsx_path, "wb") as f:
        f.write(base64.b64decode(data["content"]))

wb = openpyxl.load_workbook(xlsx_path, data_only=True)
if SHEET not in wb.sheetnames:
    raise SystemExit(f"IMPORT ABORTED: 找不到分頁「{SHEET}」，實際分頁：{wb.sheetnames[:8]}")
ws = wb[SHEET]

def find_col(pred, label):
    """依第一列標題文字找欄位（回傳 0-based index）。找不到就中止，不覆蓋既有資料。"""
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(1, c).value or "").replace("\n", "")
        if pred(h):
            return c - 1
    heads = [str(ws.cell(1, c).value or "").replace("\n", "") for c in range(1, min(ws.max_column, 25) + 1)]
    raise SystemExit(f"IMPORT ABORTED: 找不到「{label}」欄，表格結構可能變了，請人工確認。實際標題：{heads}")

COL_MEMBER = find_col(lambda h: "成員" in h and "成員收入" not in h, "成員")
COL_NAME   = find_col(lambda h: "合作名稱" in h, "合作名稱")
COL_ITEM   = find_col(lambda h: "產品項目" in h, "產品項目")
COL_START  = find_col(lambda h: "走期開始" in h, "走期開始")
COL_O      = find_col(lambda h: "成員收入" in h and "5/5" in h, "成員收入(業績<80萬 5/5分)")
COL_P      = find_col(lambda h: "成員收入" in h and "6/4" in h, "成員收入(業績>80萬 6/4分)")
COL_J      = find_col(lambda h: "成交價" in h and "未稅" in h, "成交價(未稅)")

def num(v):
    return float(v) if isinstance(v, (int, float)) else 0

out, skipped, team_rows = {}, set(), []
for r in ws.iter_rows(min_row=2):
    d = r[COL_MEMBER].value
    if d is None:
        continue
    raw = str(d).strip()
    key = MAP.get(raw)
    row = {
        "name": r[COL_NAME].value, "item": r[COL_ITEM].value,
        "start": str(r[COL_START].value)[:10] if r[COL_START].value else None,
        "o": num(r[COL_O].value), "p": num(r[COL_P].value),
    }
    team_rows.append({"member_raw": raw, "member_key": key,
                      "start": row["start"], "j": num(r[COL_J].value)})
    if not key:
        skipped.add(raw)
        continue
    out.setdefault(key, []).append(row)
for v in out.values():
    v.sort(key=lambda x: x["start"] or "9999")

total_rows = sum(len(v) for v in out.values())
if total_rows < 50:  # 目前 150+ 筆；掉到 50 以下代表來源異常，不要覆蓋
    raise SystemExit(f"IMPORT ABORTED: 只解析到 {total_rows} 筆（異常偏低），不覆蓋 revenue.json")

from datetime import datetime, timedelta, timezone
today = datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d")
json.dump({"year": 2026, "updated": today,
           "note": "來源：Google 試算表「Video 營收大表」營收大表(Video)-2026 分頁；成員頁用「成員收入」兩欄，總覽頁用「成交價(未稅)」欄（欄位以標題文字定位）",
           "by_member": out, "team_rows": team_rows},
          open(os.path.join(BASE, "revenue.json"), "w"), ensure_ascii=False, indent=1)
print(f"團隊 J 欄合計 {int(sum(t['j'] for t in team_rows)):,}（{len(team_rows)} 列）")
print(f"OK: {total_rows} 筆、{len(out)} 位成員；未列入成員名單而略過：{sorted(skipped) or '無'}")
for k in sorted(out):
    print(f"  {k}: {len(out[k])} 筆, 合計 {int(sum(x['o']+x['p'] for x in out[k])):,}")

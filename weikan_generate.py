#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""委刊單產生器：讀 spec JSON 填入 weikan_template.xlsx，輸出 xlsx + PDF。

用法：python3 weikan_generate.py spec.json [輸出目錄]

spec 範例（缺欄位就留白，客戶簽回時手填）：
{
  "no": "V20260731-01",                # 委刊單號；省略時用 V+日期-01
  "date": "2026-07-31",                # 報價日期；省略時用今天
  "version": "V1",
  "campaign_name": "和泰汽車 HINO 潮大車短影音競賽 × Rio",
  "client": {"company": "...", "address": "", "contact": "...",
              "phone": "", "email": "...", "tax_id": "", "representative": ""},
  "period": "2026/08/01 - 2026/08/31",  # 專案期間
  "member": "Rio（壹比零）",             # 合作成員（顯示在合作對象）
  "items": [{"desc": "IG 限時動態 ×1", "price": 60000}],
  "publish_date": "2026/08/19",
  "platform": "Instagram",
  "item_note": "",
  "extra_note": "",                     # 第二內容區（B13），授權範圍等
  "invoice_month": "2026 年 8 月"
}

乙方（狄卡）固定資料與匯款資訊讀 weikan_config.json。
"""
import json
import os
import subprocess
import sys
import tempfile
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parent
TEMPLATE = REPO / "weikan_template.xlsx"
CONFIG = REPO / "weikan_config.json"

AMOUNT_FMT = "#,##0"

# B11 固定尾段（沿用範本原文）
B11_TAIL = (
    "\n溝通內容、時程如雙方通訊軟體議定之內容（包含但不限於\n"
    "Email、Line、微信），若非第一次議定之內容後續額外補充，則須重新討論切角、時程、報價）\n\n"
    "- 為維護創作者風格，內容修改方式以調整資訊上錯誤為優先基礎，不提供重拍或文案大幅度修改 。"
)


def visual_width(s):
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in s)


def wrapped_lines(text, width_units):
    n = 0
    for line in text.split("\n"):
        n += max(1, -(-visual_width(line) // width_units))
    return n


def build_b11(spec):
    member = spec.get("member", "")
    lines = [f"合作對象：Dcard Video 經紀成員 {member}".rstrip(), "合作內容："]
    for it in spec.get("items", []):
        price = it.get("price")
        tag = f"　${price:,.0f}(net)" if isinstance(price, (int, float)) else ""
        lines.append(f"- {it['desc']}{tag}")
    return "\n".join(lines) + "\n" + B11_TAIL


def fill(spec, out_dir):
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))
    wb = load_workbook(TEMPLATE)
    ws = wb.active

    today = spec.get("date") or date.today().isoformat()
    client = spec.get("client", {})
    contact = cfg.get("project_contact", {})
    party_b = cfg.get("party_b", {})

    def put(coord, value, fmt=None):
        if value in (None, ""):
            return
        ws[coord] = value
        if fmt:
            ws[coord].number_format = fmt

    # 表頭
    put("B2", client.get("company"))
    put("M2", today)
    put("B3", spec.get("campaign_name"))
    put("M3", spec.get("no") or f"V{today.replace('-', '')}-01")
    put("B4", client.get("address"))
    put("M4", spec.get("version", "V1"))
    put("B5", client.get("contact"))
    put("M5", contact.get("name"))
    put("B6", client.get("phone"))
    put("M6", contact.get("mobile"))
    put("B7", client.get("email"))
    put("M7", contact.get("email"))
    put("B8", spec.get("period"))

    # B2-B8 右側被隱藏合併格(E2:K8)擋住無法溢出，值太長時縮字避免截斷
    from openpyxl.styles import Alignment
    for coord in ("B2", "B3", "B4", "B5", "B7", "B8"):
        v = ws[coord].value
        if v and visual_width(str(v)) > 34:
            a = Alignment(
                horizontal=ws[coord].alignment.horizontal,
                vertical=ws[coord].alignment.vertical,
                shrink_to_fit=True,
            )
            ws[coord].alignment = a

    # 委託項目
    b11 = build_b11(spec)
    ws["B11"] = b11
    # B11:K12 合併寬約 83 字元；依內容行數放大列高，避免 PDF 截字
    est = wrapped_lines(b11, 80)
    need = est * 14.0
    have = (ws.row_dimensions[11].height or 15) + (ws.row_dimensions[12].height or 15)
    if need > have:
        ws.row_dimensions[11].height = need - (ws.row_dimensions[12].height or 15)

    put("L11", spec.get("publish_date"))
    put("M11", spec.get("platform"))
    put("N11", spec.get("item_note"))
    put("B13", spec.get("extra_note"))

    items = spec.get("items", [])
    prices = [it.get("price") for it in items if isinstance(it.get("price"), (int, float))]
    subtotal = sum(prices) if prices else None
    if subtotal is not None:
        if len(prices) == 1:
            put("O11", prices[0], AMOUNT_FMT)
        put("P11", subtotal, AMOUNT_FMT)
        # 產出檔是簽署用文件實例，直接寫入計算結果（原範本公式 N13=P11 等保留在範本中）
        tax = round(subtotal * 0.05)
        put("N13", subtotal, AMOUNT_FMT)
        put("N15", tax, AMOUNT_FMT)
        put("N16", subtotal + tax, AMOUNT_FMT)

    # 簽章區：甲方（可留白由客戶手填）
    put("B32", client.get("company"))
    put("B33", client.get("address"))
    put("B34", client.get("phone"))
    put("B35", client.get("tax_id"))
    put("B36", client.get("representative"))
    # 乙方（狄卡）
    put("L32", party_b.get("company"))
    put("L33", party_b.get("address"))
    put("L34", party_b.get("phone"))
    put("L35", party_b.get("tax_id"))
    put("L36", party_b.get("contact") or contact.get("name"))
    put("M39", cfg.get("remittance"))
    put("B42", spec.get("invoice_month"))

    # 版面：單頁寬、直式 A4
    ws.print_area = "A1:Q43"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = "".join(c for c in spec.get("campaign_name", "委刊單") if c not in '\\/:*?"<>|')[:60]
    xlsx_path = out_dir / f"委刊單_{safe}.xlsx"
    wb.save(xlsx_path)
    return xlsx_path


def _soffice_convert(xlsx_path, out_dir):
    # sandbox 環境 soffice 需要獨立 user profile（不然會載入失敗）
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    with tempfile.TemporaryDirectory(prefix="lo_profile_") as profile:
        return subprocess.run(
            ["soffice", f"-env:UserInstallation={Path(profile).as_uri()}",
             "--headless", "--convert-to", "pdf", "--outdir", out_dir, str(xlsx_path)],
            capture_output=True, text=True, timeout=180, env=env,
        )


def to_pdf(xlsx_path):
    out_dir = str(Path(xlsx_path).parent)
    r = _soffice_convert(xlsx_path, out_dir)
    pdf_path = Path(xlsx_path).with_suffix(".pdf")
    if not pdf_path.exists() and "could not be loaded" in (r.stderr or ""):
        # 容器重建後只剩 libreoffice-core，補裝 calc 模組再重試一次
        print("libreoffice-calc 缺失，嘗試安裝……", file=sys.stderr)
        subprocess.run(
            "apt-get update -qq && apt-get install -y -qq --no-install-recommends libreoffice-calc",
            shell=True, capture_output=True, text=True, timeout=600,
        )
        r = _soffice_convert(xlsx_path, out_dir)
    if not pdf_path.exists():
        raise RuntimeError(f"PDF 轉檔失敗：{r.stdout}\n{r.stderr}")
    return pdf_path


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out_dir = sys.argv[2] if len(sys.argv) > 2 else str(REPO / "weikan_out")
    xlsx_path = fill(spec, out_dir)
    pdf_path = to_pdf(xlsx_path)
    print(json.dumps({"xlsx": str(xlsx_path), "pdf": str(pdf_path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
